from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

from django.db.models import Q

from app.models.classes import Class, Term
from app.models.fees_payment import (
    Payment,
    StudentBillItem,
    infer_ledger_category,
    normalize_payment_method,
)
from app.models.school_settings import AcademicYear
from app.models.students import Student

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except Exception:
    Workbook = None
    Alignment = Font = PatternFill = None


def parse_date_value(raw_value):
    value = str(raw_value or "").strip()
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return None


def get_student_reg_no(student):
    # The ledger must expose the client's standardized numeric Student ID when available.
    return getattr(student, "student_number", "") or getattr(student, "reg_no", "") or str(student.pk)


def get_classroom_label(academic_class):
    if not academic_class or not getattr(academic_class, "Class", None):
        return "-"
    return getattr(academic_class.Class, "code", None) or getattr(academic_class.Class, "name", None) or str(academic_class.Class)


def get_term_label(academic_class):
    if not academic_class:
        return "-"
    term_number = getattr(getattr(academic_class, "term", None), "term", None)
    academic_year = getattr(getattr(academic_class, "academic_year", None), "academic_year", None)
    if term_number and academic_year:
        return f"Term {term_number} {academic_year}"
    if term_number:
        return f"Term {term_number}"
    return "-"


def get_classroom_group(classroom_label):
    normalized = str(classroom_label or "").strip().upper()
    if normalized.startswith("KG"):
        return "KG"
    if normalized.startswith("GD"):
        return "GD"
    if "TAHF" in normalized:
        return "TAHFIDH"
    return ""


def get_balance_status(balance):
    if balance < 0:
        return "credit"
    if balance > 0:
        return "outstanding"
    return "cleared"


def get_balance_status_label(balance):
    return get_balance_status(balance).title()


def _apply_common_filters(queryset, *, student_id="", class_id="", academic_year_id="", term_id="", classroom_group="", classroom_code=""):
    if student_id:
        queryset = queryset.filter(bill__student_id=student_id)
    if class_id:
        queryset = queryset.filter(bill__academic_class__Class_id=class_id)
    if classroom_code:
        queryset = queryset.filter(
            Q(bill__academic_class__Class__code__iexact=classroom_code)
            | Q(bill__academic_class__Class__name__iexact=classroom_code)
        )
    if academic_year_id:
        queryset = queryset.filter(bill__academic_class__academic_year_id=academic_year_id)
    if term_id:
        queryset = queryset.filter(bill__academic_class__term_id=term_id)
    if classroom_group == "KG":
        queryset = queryset.filter(
            Q(bill__academic_class__Class__code__istartswith="KG")
            | Q(bill__academic_class__Class__name__istartswith="KG")
        )
    elif classroom_group == "GD":
        queryset = queryset.filter(
            Q(bill__academic_class__Class__code__istartswith="GD")
            | Q(bill__academic_class__Class__name__istartswith="GD")
        )
    elif classroom_group == "TAHFIDH":
        queryset = queryset.filter(
            Q(bill__academic_class__Class__code__icontains="TAHF")
            | Q(bill__academic_class__Class__name__icontains="TAHF")
        )
    return queryset


def _resolve_payment_category(payment):
    if getattr(payment, "fee_category", ""):
        return payment.fee_category

    item_categories = [
        category for category in payment.bill.items.values_list("fee_category", flat=True).distinct() if category
    ]
    if len(item_categories) == 1:
        return item_categories[0]

    return infer_ledger_category(*payment.bill.items.values_list("description", flat=True))


def build_ledger_rows(
    *,
    student_id="",
    class_id="",
    academic_year_id="",
    term_id="",
    classroom_group="",
    classroom_code="",
    categories=None,
    payment_methods=None,
    date_from=None,
    date_to=None,
    balance_mode="student",
):
    categories = [value for value in (categories or []) if value]
    payment_methods = [normalize_payment_method(value) for value in (payment_methods or []) if value]
    if balance_mode not in {"overall", "student"}:
        balance_mode = "student"

    charge_qs = StudentBillItem.objects.select_related(
        "bill",
        "bill__student",
        "bill__academic_class",
        "bill__academic_class__Class",
        "bill__academic_class__term",
        "bill__academic_class__academic_year",
        "bill_item",
    )
    payment_qs = Payment.objects.select_related(
        "bill",
        "bill__student",
        "bill__academic_class",
        "bill__academic_class__Class",
        "bill__academic_class__term",
        "bill__academic_class__academic_year",
    )

    charge_qs = _apply_common_filters(
        charge_qs,
        student_id=student_id,
        class_id=class_id,
        academic_year_id=academic_year_id,
        term_id=term_id,
        classroom_group=classroom_group,
        classroom_code=classroom_code,
    )
    payment_qs = _apply_common_filters(
        payment_qs,
        student_id=student_id,
        class_id=class_id,
        academic_year_id=academic_year_id,
        term_id=term_id,
        classroom_group=classroom_group,
        classroom_code=classroom_code,
    )

    if categories:
        charge_qs = charge_qs.filter(fee_category__in=categories)

    if payment_methods:
        payment_qs = payment_qs.filter(payment_method__in=payment_methods)

    if date_from:
        charge_qs = charge_qs.filter(charge_date__gte=date_from)
        payment_qs = payment_qs.filter(payment_date__gte=date_from)
    if date_to:
        charge_qs = charge_qs.filter(charge_date__lte=date_to)
        payment_qs = payment_qs.filter(payment_date__lte=date_to)

    rows = []

    for charge in charge_qs.order_by("charge_date", "id"):
        charge_date = charge.charge_date or charge.bill.bill_date
        classroom = get_classroom_label(charge.bill.academic_class)
        charge_amount = Decimal(charge.amount or 0)
        is_credit_adjustment = charge_amount < 0
        description = charge.description or charge.notes or getattr(charge.bill_item, "item_name", "")
        notes = charge.notes or charge.description or getattr(charge.bill_item, "description", "") or ""
        rows.append(
            {
                "row_type": "charge",
                "bill_id": charge.bill_id,
                "student_pk": charge.bill.student_id,
                "sort_id": charge.id,
                "date": charge_date,
                "reg_no": get_student_reg_no(charge.bill.student),
                "student_name": charge.bill.student.student_name,
                "classroom": classroom,
                "classroom_group": get_classroom_group(classroom),
                "term": get_term_label(charge.bill.academic_class),
                "type_label": "Credit Adjustment" if is_credit_adjustment else "Charge",
                "type_badge_class": "credit-adjustment" if is_credit_adjustment else "charge",
                "category": charge.fee_category or infer_ledger_category(
                    getattr(charge.bill_item, "category", ""),
                    getattr(charge.bill_item, "item_name", ""),
                    charge.description,
                ),
                "payment_method": "",
                "amount_charged": charge_amount,
                "amount_charged_display": abs(charge_amount),
                "amount_paid": Decimal("0"),
                "amount_paid_display": Decimal("0"),
                "reference": f"CHG-{charge.bill_id}-{charge.id}",
                "description": description or ("Credit adjustment" if is_credit_adjustment else "-"),
                "notes": notes,
            }
        )

    for payment in payment_qs.order_by("payment_date", "id"):
        classroom = get_classroom_label(payment.bill.academic_class)
        category = _resolve_payment_category(payment)
        paid_amount = Decimal(payment.amount or 0)
        if categories and category not in categories:
            continue
        rows.append(
            {
                "row_type": "payment",
                "bill_id": payment.bill_id,
                "student_pk": payment.bill.student_id,
                "sort_id": payment.id,
                "date": payment.payment_date,
                "reg_no": get_student_reg_no(payment.bill.student),
                "student_name": payment.bill.student.student_name,
                "classroom": classroom,
                "classroom_group": get_classroom_group(classroom),
                "term": get_term_label(payment.bill.academic_class),
                "type_label": "Payment",
                "type_badge_class": "payment",
                "category": category,
                "payment_method": normalize_payment_method(payment.payment_method),
                "amount_charged": Decimal("0"),
                "amount_charged_display": Decimal("0"),
                "amount_paid": paid_amount,
                "amount_paid_display": paid_amount,
                "reference": payment.reference_no,
                "description": payment.notes or "Payment received",
                "notes": payment.notes or "",
            }
        )

    if balance_mode == "student":
        rows.sort(
            key=lambda row: (
                str(row["student_name"]).lower(),
                str(row["reg_no"]),
                row["date"],
                0 if row["row_type"] == "charge" else 1,
                row["sort_id"],
            )
        )
    else:
        rows.sort(
            key=lambda row: (
                row["date"],
                0 if row["row_type"] == "charge" else 1,
                row["sort_id"],
            )
        )

    running_balance = Decimal("0")
    student_balances = {}
    total_charged = Decimal("0")
    total_paid = Decimal("0")
    for row in rows:
        total_charged += row["amount_charged"]
        total_paid += row["amount_paid"]
        if balance_mode == "student":
            balance_key = row["student_pk"]
            student_balances.setdefault(balance_key, Decimal("0"))
            student_balances[balance_key] += row["amount_charged"]
            student_balances[balance_key] -= row["amount_paid"]
            row["running_balance"] = student_balances[balance_key]
        else:
            running_balance += row["amount_charged"]
            running_balance -= row["amount_paid"]
            row["running_balance"] = running_balance
        row["balance_status"] = get_balance_status(row["running_balance"])
        row["balance_status_label"] = get_balance_status_label(row["running_balance"])

    if balance_mode == "student":
        running_balance = sum(student_balances.values(), Decimal("0"))

    student_summaries = {}
    positive_charge_category_totals = defaultdict(lambda: Decimal("0"))
    absolute_charge_category_totals = defaultdict(lambda: Decimal("0"))
    positive_charge_total = Decimal("0")

    for row in rows:
        summary = student_summaries.setdefault(
            row["student_pk"],
            {
                "student_pk": row["student_pk"],
                "student_name": row["student_name"],
                "reg_no": row["reg_no"],
                "classroom": row["classroom"],
                "term": row["term"],
                "total_charged": Decimal("0"),
                "total_paid": Decimal("0"),
                "balance": Decimal("0"),
                "balance_abs": Decimal("0"),
                "transaction_count": 0,
                "latest_activity_date": row["date"],
                "latest_activity_sort": None,
                "action_bill_id": row["bill_id"],
            },
        )
        summary["total_charged"] += row["amount_charged"]
        summary["total_paid"] += row["amount_paid"]
        summary["balance"] = summary["total_charged"] - summary["total_paid"]
        summary["balance_abs"] = abs(summary["balance"])
        summary["transaction_count"] += 1

        latest_activity_sort = (
            row["date"],
            0 if row["row_type"] == "charge" else 1,
            row["sort_id"],
        )
        if summary["latest_activity_sort"] is None or latest_activity_sort > summary["latest_activity_sort"]:
            summary["latest_activity_sort"] = latest_activity_sort
            summary["latest_activity_date"] = row["date"]
            summary["action_bill_id"] = row["bill_id"]
            summary["classroom"] = row["classroom"]
            summary["term"] = row["term"]

        if row["row_type"] == "charge":
            category = row["category"] or "Other"
            absolute_charge_category_totals[category] += abs(row["amount_charged"])
            if row["amount_charged"] > 0:
                positive_charge_category_totals[category] += row["amount_charged"]
                positive_charge_total += row["amount_charged"]

    ordered_student_summaries = sorted(
        student_summaries.values(),
        key=lambda summary: (
            str(summary["student_name"]).lower(),
            str(summary["reg_no"]),
        ),
    )
    for summary in ordered_student_summaries:
        summary["balance_status"] = get_balance_status(summary["balance"])
        summary["balance_status_label"] = get_balance_status_label(summary["balance"])

    total_outstanding = sum(
        (summary["balance"] for summary in ordered_student_summaries if summary["balance"] > 0),
        Decimal("0"),
    )
    total_credit = sum(
        (summary["balance_abs"] for summary in ordered_student_summaries if summary["balance"] < 0),
        Decimal("0"),
    )
    top_outstanding_student = max(
        (summary for summary in ordered_student_summaries if summary["balance"] > 0),
        key=lambda summary: summary["balance"],
        default=None,
    )
    top_credit_student = max(
        (summary for summary in ordered_student_summaries if summary["balance"] < 0),
        key=lambda summary: summary["balance_abs"],
        default=None,
    )

    charge_category_totals = positive_charge_category_totals if positive_charge_total > 0 else absolute_charge_category_totals
    charge_total_for_share = positive_charge_total if positive_charge_total > 0 else sum(
        absolute_charge_category_totals.values(),
        Decimal("0"),
    )
    largest_charge_category = None
    if charge_category_totals:
        category_label, category_amount = max(
            charge_category_totals.items(),
            key=lambda item: item[1],
        )
        largest_charge_category = {
            "label": category_label,
            "amount": category_amount,
            "share": ((category_amount / charge_total_for_share) * Decimal("100")) if charge_total_for_share else Decimal("0"),
        }

    for row in rows:
        summary = student_summaries[row["student_pk"]]
        row["student_total_charged"] = summary["total_charged"]
        row["student_total_paid"] = summary["total_paid"]
        row["student_balance"] = summary["balance"]
        row["student_balance_status"] = summary["balance_status"]
        row["student_balance_status_label"] = summary["balance_status_label"]
        row["student_action_bill_id"] = summary["action_bill_id"]
        row["student_transaction_count"] = summary["transaction_count"]
        row["student_latest_activity"] = summary["latest_activity_date"]

    return {
        "rows": rows,
        "total_charged": total_charged,
        "total_paid": total_paid,
        "running_balance": running_balance,
        "transaction_count": len(rows),
        "student_count": len(ordered_student_summaries),
        "credit_row_count": sum(1 for row in rows if row["running_balance"] < 0),
        "student_summaries": ordered_student_summaries,
        "total_outstanding": total_outstanding,
        "total_credit": total_credit,
        "top_outstanding_student": top_outstanding_student,
        "top_credit_student": top_credit_student,
        "largest_charge_category": largest_charge_category,
    }


def get_ledger_filter_options():
    return {
        "students": Student.objects.order_by("student_name", "student_number", "reg_no"),
        "classes": Class.objects.order_by("code", "name"),
        "academic_years": AcademicYear.objects.order_by("-academic_year"),
        "terms": Term.objects.select_related("academic_year").order_by("-academic_year__academic_year", "term"),
    }


def _column_letter(column_number):
    result = []
    while column_number:
        column_number, remainder = divmod(column_number - 1, 26)
        result.append(chr(65 + remainder))
    return "".join(reversed(result))


def _inline_string_cell(cell_ref, value):
    safe_value = escape(str(value or ""))
    if safe_value.startswith(" ") or safe_value.endswith(" "):
        return f'<c r="{cell_ref}" t="inlineStr"><is><t xml:space="preserve">{safe_value}</t></is></c>'
    return f'<c r="{cell_ref}" t="inlineStr"><is><t>{safe_value}</t></is></c>'


def _number_cell(cell_ref, value):
    return f'<c r="{cell_ref}"><v>{value}</v></c>'


def _build_fallback_workbook_bytes(rows):
    headers = [
        "Date",
        "Student ID",
        "Student Name",
        "Classroom",
        "Term",
        "Category",
        "Payment Method",
        "Amount Charged",
        "Amount Paid",
        "Notes",
    ]

    row_xml = []
    for row_index, header_row in enumerate([headers], start=1):
        cells = [
            _inline_string_cell(f"{_column_letter(column_index)}{row_index}", value)
            for column_index, value in enumerate(header_row, start=1)
        ]
        row_xml.append(f'<row r="{row_index}">{"".join(cells)}</row>')

    for row_index, row in enumerate(rows, start=2):
        values = [
            row["date"].strftime("%Y-%m-%d") if getattr(row["date"], "strftime", None) else str(row["date"] or ""),
            row["reg_no"],
            row["student_name"],
            row["classroom"],
            row["term"],
            row["category"],
            row["payment_method"],
            row["amount_charged"],
            row["amount_paid"],
            row["notes"],
        ]
        cells = []
        for column_index, value in enumerate(values, start=1):
            cell_ref = f"{_column_letter(column_index)}{row_index}"
            if column_index in {8, 9}:
                cells.append(_number_cell(cell_ref, value))
            else:
                cells.append(_inline_string_cell(cell_ref, value))
        row_xml.append(f'<row r="{row_index}">{"".join(cells)}</row>')

    worksheet_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetData>'
        f'{"".join(row_xml)}'
        '</sheetData>'
        '</worksheet>'
    )

    workbook_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Ledger" sheetId="1" r:id="rId1"/></sheets>'
        '</workbook>'
    )
    workbook_rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
        'Target="worksheets/sheet1.xml"/>'
        '</Relationships>'
    )
    root_rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="xl/workbook.xml"/>'
        '</Relationships>'
    )
    content_types_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '</Types>'
    )

    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types_xml)
        archive.writestr("_rels/.rels", root_rels_xml)
        archive.writestr("xl/workbook.xml", workbook_xml)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_rels_xml)
        archive.writestr("xl/worksheets/sheet1.xml", worksheet_xml)
    return output.getvalue()


def build_ledger_workbook(rows):
    if Workbook is None:
        return _build_fallback_workbook_bytes(rows)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ledger"

    headers = [
        "Date",
        "Student ID",
        "Student Name",
        "Classroom",
        "Term",
        "Category",
        "Payment Method",
        "Amount Charged",
        "Amount Paid",
        "Notes",
    ]
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        sheet.append(
            [
                row["date"].strftime("%Y-%m-%d") if getattr(row["date"], "strftime", None) else str(row["date"] or ""),
                row["reg_no"],
                row["student_name"],
                row["classroom"],
                row["term"],
                row["category"],
                row["payment_method"],
                float(row["amount_charged"]),
                float(row["amount_paid"]),
                row["notes"],
            ]
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
